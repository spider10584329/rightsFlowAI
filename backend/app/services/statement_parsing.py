"""Statement file parsing and the import adapter registry.

An adapter knows the column vocabulary of one statement family and proposes a
mapping from file headers to canonical fields. New formats are added by
registering another ``StatementAdapter`` in ``ADAPTERS``; nothing else changes.
"""

import csv
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from app.core.errors import ValidationFailedError

MAX_ROWS = 250_000
MAX_COLUMNS = 200
MAX_BLANK_ROWS = 10_000

CANONICAL_FIELDS: dict[str, str] = {
    "isrc": "ISRC (recording identifier)",
    "iswc": "ISWC (work identifier)",
    "upc": "UPC / EAN (release identifier)",
    "title": "Track or work title",
    "artist": "Artist",
    "territory": "Territory / country of sale",
    "usage_type": "Usage / transaction type",
    "rights_type": "Rights type",
    "store": "Store / service",
    "quantity": "Units / streams",
    "gross_amount": "Gross amount",
    "amount": "Net amount payable to you (required)",
    "currency": "Currency",
    "transaction_date": "Transaction / sale date",
    "source_reference": "Source line reference",
}
IDENTIFYING_FIELDS = ("isrc", "iswc", "upc", "title")


def _norm(header: str) -> str:
    return re.sub(r"[^a-z0-9]", "", header.lower())


@dataclass
class StatementAdapter:
    key: str
    label: str
    description: str
    aliases: dict[str, list[str]]
    signature_headers: list[str] = field(default_factory=list)

    def suggest_mapping(self, headers: list[str]) -> dict[str, str]:
        normalized = {_norm(h): h for h in headers}
        mapping: dict[str, str] = {}
        used: set[str] = set()
        for canonical, aliases in self.aliases.items():
            for alias in [canonical, *aliases]:
                header = normalized.get(_norm(alias))
                if header and header not in used:
                    mapping[canonical] = header
                    used.add(header)
                    break
        return mapping

    def detect_score(self, headers: list[str]) -> float:
        normalized = {_norm(h) for h in headers}
        mapped = len(self.suggest_mapping(headers))
        signature = sum(1 for s in self.signature_headers if _norm(s) in normalized)
        return mapped + 2 * signature


_COMMON = {
    "isrc": ["isrc code", "track isrc", "recording isrc"],
    "iswc": ["iswc code", "work iswc"],
    "upc": ["ean", "upc code", "barcode", "product upc", "release upc"],
    "title": ["track title", "track name", "song title", "song", "track", "work title",
              "title name", "product title"],
    "artist": ["artist name", "performer", "track artist", "main artist"],
    "territory": ["country", "country code", "country of sale", "territory code", "region",
                  "sale country"],
    "usage_type": ["usage", "sale type", "transaction type", "type", "use type",
                   "income type", "revenue type"],
    "rights_type": ["right type", "royalty type", "rights"],
    "store": ["service", "platform", "dsp", "retailer", "store name", "source"],
    "quantity": ["units", "qty", "streams", "plays", "downloads", "performances", "count"],
    "gross_amount": ["gross", "gross revenue", "gross amount", "gross earnings"],
    "amount": ["net amount", "net revenue", "earnings", "royalty", "royalty amount", "amount due",
               "net", "payable", "revenue", "earnings usd", "total"],
    "currency": ["currency code", "ccy", "reporting currency"],
    "transaction_date": ["date", "sale date", "reporting date", "period", "sale month",
                         "statement period", "usage date"],
    "source_reference": ["line id", "reference", "transaction id", "row id"],
}


def _merge(extra: dict[str, list[str]]) -> dict[str, list[str]]:
    merged = {k: list(v) for k, v in _COMMON.items()}
    for key, aliases in extra.items():
        merged[key] = aliases + merged.get(key, [])
    return merged


ADAPTERS: dict[str, StatementAdapter] = {
    adapter.key: adapter
    for adapter in [
        StatementAdapter(
            key="generic",
            label="Generic CSV / XLSX",
            description="Any tabular statement; map the columns manually if needed.",
            aliases=_merge({}),
        ),
        StatementAdapter(
            key="distributor_standard",
            label="Digital distributor sales report",
            description="Store-level sales/streams with ISRC, UPC, country and earnings.",
            aliases=_merge({
                "amount": ["earnings (usd)", "earnings usd", "net earnings", "royalty (usd)"],
                "territory": ["country of sale", "sale country"],
                "store": ["store", "retailer"],
                "transaction_date": ["sale month", "reporting month", "sale period"],
            }),
            signature_headers=["store", "country of sale", "earnings (usd)", "sale month", "upc"],
        ),
        StatementAdapter(
            key="dsp_streaming",
            label="Streaming platform report",
            description="DSP-style stream counts and revenue per track and territory.",
            aliases=_merge({
                "quantity": ["streams", "stream count", "plays"],
                "amount": ["revenue", "net revenue", "payable amount"],
                "title": ["track name", "track title"],
            }),
            signature_headers=["streams", "track name", "artist name", "revenue"],
        ),
        StatementAdapter(
            key="pro_works",
            label="PRO / publishing works statement",
            description="Work-level performance or mechanical royalties keyed by ISWC.",
            aliases=_merge({
                "title": ["work title", "composition title", "title"],
                "usage_type": ["performance type", "use type", "source of income", "income type"],
                "amount": ["royalty amount", "amount", "writer share", "publisher share",
                           "distribution amount"],
                "transaction_date": ["performance date", "distribution date", "period"],
            }),
            signature_headers=["work title", "iswc", "performance type", "royalty amount"],
        ),
    ]
}


def detect_adapter(headers: list[str]) -> StatementAdapter:
    best = max(ADAPTERS.values(), key=lambda a: (a.detect_score(headers), a.key != "generic"))
    return best


# --------------------------------------------------------------------------- file reading


@dataclass
class ParsedTable:
    headers: list[str]
    rows: list[dict[str, str]]
    delimiter: str | None = None  # CSV only; informs decimal-separator detection


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ValidationFailedError("Unable to decode the statement file")


def _dedupe_headers(raw: list[str]) -> list[str]:
    headers: list[str] = []
    seen: dict[str, int] = {}
    for index, value in enumerate(raw):
        name = (value or "").strip() or f"Column {index + 1}"
        if name in seen:
            seen[name] += 1
            name = f"{name} ({seen[name]})"
        else:
            seen[name] = 1
        headers.append(name)
    return headers


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float):
        # Excel stores numbers as binary floats; repr gives the shortest
        # round-tripping decimal string, avoiding 0.1+0.2 artefacts.
        return repr(value)
    return str(value).strip()


def read_table(data: bytes, extension: str, max_rows: int = MAX_ROWS) -> ParsedTable:
    extension = extension.lower()
    if extension == ".xlsx":
        return _read_xlsx(data, max_rows)
    if extension in (".csv", ".tsv", ".txt"):
        return _read_csv(data, max_rows, "\t" if extension == ".tsv" else None)
    raise ValidationFailedError(f"Unsupported statement format {extension}")


def _read_csv(data: bytes, max_rows: int, delimiter: str | None) -> ParsedTable:
    text = _decode(data)
    if delimiter is None:
        sample = text[:20000]
        try:
            delimiter = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
        except csv.Error:
            delimiter = ","
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    try:
        raw_rows = [row for row in reader if any(cell.strip() for cell in row)]
    except csv.Error as exc:  # e.g. a field over the 128 KB limit (audit AF-11)
        raise ValidationFailedError(f"The CSV file is malformed: {exc}") from exc
    if not raw_rows:
        raise ValidationFailedError("The statement file has no rows")
    headers = _dedupe_headers(raw_rows[0])
    body = raw_rows[1:]
    if len(body) > max_rows:
        raise ValidationFailedError(f"Statements are limited to {max_rows} rows per file")
    rows = [
        {headers[i]: (row[i].strip() if i < len(row) else "") for i in range(len(headers))}
        for row in body
    ]
    return ParsedTable(headers=headers, rows=rows, delimiter=delimiter)


def _read_xlsx(data: bytes, max_rows: int) -> ParsedTable:
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises many types for corrupt files
        raise ValidationFailedError("Unable to read the XLSX workbook") from exc
    try:
        sheet = workbook.worksheets[0]
        # Don't trust the sheet's declared <dimension>: a crafted file can claim
        # 1M rows x 16k columns with two real cells (audit AF-01/02).
        sheet.reset_dimensions()
        iterator = sheet.iter_rows(values_only=True, max_col=MAX_COLUMNS)
        headers: list[str] | None = None
        rows: list[dict[str, str]] = []
        scanned = 0
        for raw in iterator:
            scanned += 1
            if scanned > max_rows + MAX_BLANK_ROWS:
                raise ValidationFailedError(f"Statements are limited to {max_rows} rows per file")
            cells = [_cell_text(v) for v in raw]
            if not any(cells):
                continue
            if headers is None:
                headers = _dedupe_headers(cells)
                continue
            if len(rows) >= max_rows:
                raise ValidationFailedError(f"Statements are limited to {max_rows} rows per file")
            rows.append({headers[i]: (cells[i] if i < len(cells) else "") for i in range(len(headers))})
    finally:
        workbook.close()
    if headers is None:
        raise ValidationFailedError("The workbook has no rows")
    return ParsedTable(headers=headers, rows=rows)


# --------------------------------------------------------------------------- value parsing

_SPACES = re.compile(r"[\s   ']")
_CURRENCY_SYMBOLS = "$€£¥₩₹"
_EDGE_CURRENCY = re.compile(r"^(?:[A-Za-z]{3}|[$€£¥₩₹])+|(?:[A-Za-z]{3}|[$€£¥₩₹])+$")
_NUMBER_CHARS = re.compile(r"^[0-9.,]+([eE][+-]?\d+)?$")


def _strip_number(value: str) -> tuple[str, bool]:
    """Remove spaces, sign wrappers and currency markers at the edges.

    Returns (digits-and-separators, negative). Raises ValueError for anything
    that is not clearly a number: letters may only appear as a leading or
    trailing currency code/symbol (e.g. ``USD 12.50``, ``12,50 €``).
    """
    text = _SPACES.sub("", value)
    negative = False
    if text.startswith("(") and text.endswith(")"):
        negative, text = True, text[1:-1]
        if text.startswith(("-", "+")):
            raise ValueError(f"Not a number: {value!r}")
    text = _EDGE_CURRENCY.sub("", text)
    if text.endswith("-"):
        negative, text = not negative, text[:-1]
    if text.startswith("-"):
        negative, text = not negative, text[1:]
    elif text.startswith("+"):
        text = text[1:]
    text = _EDGE_CURRENCY.sub("", text)
    if not text or not _NUMBER_CHARS.match(text):
        raise ValueError(f"Not a number: {value!r}")
    return text, negative


def separator_evidence(value: str) -> str | None:
    """Which decimal separator this single value proves, if any."""
    try:
        text, _ = _strip_number(value)
    except ValueError:
        return None
    mantissa = re.split(r"[eE]", text)[0]
    commas, dots = mantissa.count(","), mantissa.count(".")
    if commas and dots:
        return "," if mantissa.rfind(",") > mantissa.rfind(".") else "."
    if commas:
        whole, _, frac = mantissa.rpartition(",")
        if commas > 1:
            return "."  # 1,234,567 — commas group thousands
        if len(frac) != 3 or whole in ("", "0"):
            return ","  # 12,5 / 0,004 — a decimal comma
        return None  # 1,234 is ambiguous on its own
    if dots:
        whole, _, frac = mantissa.rpartition(".")
        if dots > 1:
            return ","  # 1.234.567 — dots group thousands
        if len(frac) != 3 or whole in ("", "0"):
            return "."
        return None
    return None


def detect_decimal_separator(values: list[str | None], csv_delimiter: str | None = None) -> str:
    """Decide one decimal separator for a whole column from the values that
    prove it; ambiguous columns default to "," for ';'-delimited files (the
    European convention) and "." otherwise."""
    votes = {",": 0, ".": 0}
    for value in values:
        if value:
            evidence = separator_evidence(value)
            if evidence:
                votes[evidence] += 1
    if votes[","] != votes["."]:
        return "," if votes[","] > votes["."] else "."
    return "," if csv_delimiter == ";" else "."


def parse_decimal(value: str | None, decimal_separator: str | None = None) -> Decimal | None:
    """Parse amounts like ``1,234.56``, ``1.234,56``, ``(12.50)``, ``-3``, ``$9``.

    ``decimal_separator`` should come from :func:`detect_decimal_separator`
    for the column; without it the value's own evidence is used and ambiguous
    groups like ``1,234`` are read with "." as the decimal separator.
    Returns None for blanks; raises ValueError for unparseable text.
    """
    if value is None or not value.strip():
        return None
    text, negative = _strip_number(value)
    separator = decimal_separator or separator_evidence(value) or "."
    thousands = "," if separator == "." else "."
    mantissa, exponent = text, ""
    if re.search(r"[eE]", text):
        mantissa, power = re.split(r"[eE]", text, maxsplit=1)
        exponent = "e" + power
    if mantissa.count(separator) > 1:
        raise ValueError(f"Not a number: {value!r}")
    mantissa = mantissa.replace(thousands, "").replace(separator, ".")
    if not re.match(r"^(\d+\.?\d*|\.\d+)$", mantissa):
        raise ValueError(f"Not a number: {value!r}")
    try:
        number = Decimal(mantissa + exponent)
    except InvalidOperation as exc:
        raise ValueError(f"Not a number: {value!r}") from exc
    if not number.is_finite():
        raise ValueError(f"Not a number: {value!r}")
    return -number if negative else number


_MONTH_NAME_FORMATS = ("%b %Y", "%B %Y", "%b-%y", "%b-%Y", "%B-%Y", "%d %b %Y", "%d %B %Y", "%b %d, %Y")
_YEAR_LAST = re.compile(r"^(\d{1,2})[/.\-](\d{1,2})[/.\-](\d{4})$")


def detect_date_order(values: list[str | None]) -> str:
    """'dmy' or 'mdy' for day/month/year dates in a column (e.g. 03/04/2026).
    A day > 12 in either position decides; dot-separated dates and otherwise
    ambiguous columns default to day-first, except purely slash-separated
    ambiguous columns, which default to month-first (US distributor exports)."""
    dmy = mdy = 0
    dots = slashes = 0
    for value in values:
        if not value:
            continue
        match = _YEAR_LAST.match(value.strip())
        if not match:
            continue
        first, second = int(match.group(1)), int(match.group(2))
        dots += "." in value
        slashes += "/" in value
        if first > 12 >= second:
            dmy += 1
        elif second > 12 >= first:
            mdy += 1
    if dmy != mdy:
        return "dmy" if dmy > mdy else "mdy"
    return "mdy" if slashes and not dots else "dmy"


def parse_date(value: str | None, order: str = "dmy") -> date | None:
    """Parse ISO dates, compact ``YYYYMMDD`` / ``YYYYMM``, ``YYYY-MM``,
    day/month/year in the column's ``order``, and month-name forms.
    Raises ValueError for anything unrecognized."""
    if value is None or not value.strip():
        return None
    text = value.strip()
    if "T" in text:
        text = text.split("T", 1)[0]
    try:
        if re.fullmatch(r"\d{8}", text):
            return datetime.strptime(text, "%Y%m%d").date()
        if re.fullmatch(r"\d{6}", text):
            return datetime.strptime(text, "%Y%m").date()
        if re.fullmatch(r"\d{4}[-/.]\d{1,2}[-/.]\d{1,2}", text):
            year, month, day = re.split(r"[-/.]", text)
            return date(int(year), int(month), int(day))
        if re.fullmatch(r"\d{4}[-/.]\d{1,2}", text):
            year, month = re.split(r"[-/.]", text)
            return date(int(year), int(month), 1)
        match = _YEAR_LAST.match(text)
        if match:
            first, second, year = (int(g) for g in match.groups())
            day, month = (first, second) if order == "dmy" else (second, first)
            return date(year, month, day)
        if re.fullmatch(r"\d{1,2}/\d{4}", text):
            month, year = text.split("/")
            return date(int(year), int(month), 1)
    except ValueError as exc:
        raise ValueError(f"Unrecognized date: {value!r}") from exc
    for fmt in _MONTH_NAME_FORMATS:
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"Unrecognized date: {value!r}")


_RIGHTS_WORDS = {
    "perf": "performance", "performance": "performance", "public performance": "performance",
    "mech": "mechanical", "mechanical": "mechanical", "mechanicals": "mechanical",
    "sync": "sync", "synchronization": "sync", "synchronisation": "sync",
    "print": "print", "lyrics": "print",
    "master": "master", "sound recording": "master", "recording": "master",
    "neighbouring": "neighbouring", "neighboring": "neighbouring", "nr": "neighbouring",
}


def parse_rights_type(value: str | None) -> str | None:
    if not value:
        return None
    return _RIGHTS_WORDS.get(value.strip().lower())


_PLAIN_NUMBER_RE = re.compile(r"[+-]?\d+(?:\.\d+)?(?:[eE][+-]?\d+)?")


def neutralize_formula(value: str) -> str:
    """Protect CSV exports against spreadsheet formula injection.

    Plain numbers such as ``-12.50`` are left alone: prefixing them turned
    every negative amount into text (audit F-M8). Anything else starting
    with a formula trigger is prefixed with an apostrophe."""
    if value and value[0] in ("=", "+", "-", "@", "\t", "\r"):
        if _PLAIN_NUMBER_RE.fullmatch(value):
            return value
        return "'" + value
    return value
